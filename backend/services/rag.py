"""RAG (Retrieval-Augmented Generation) pipeline."""

import io
import uuid

import PyPDF2
import docx
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from core.config import settings
from models.database import Document, DocumentChunk
from services.llm import get_embedding


def chunk_text(text: str, chunk_size: int = None, overlap: int = None) -> list[str]:
    chunk_size = chunk_size or settings.chunk_size
    overlap = overlap or settings.chunk_overlap

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk.strip())
        start = end - overlap
        if start >= len(text):
            break
    return chunks


def extract_text_from_pdf(content: bytes) -> str:
    reader = PyPDF2.PdfReader(io.BytesIO(content))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text_from_docx(content: bytes) -> str:
    doc = docx.Document(io.BytesIO(content))
    return "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())


def extract_text_from_txt(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


def extract_text(content: bytes, content_type: str) -> str:
    if "pdf" in content_type:
        return extract_text_from_pdf(content)
    elif "word" in content_type or "docx" in content_type:
        return extract_text_from_docx(content)
    else:
        return extract_text_from_txt(content)


async def ingest_document(
    db: AsyncSession,
    filename: str,
    content_type: str,
    file_content: bytes,
) -> Document:
    raw_text = extract_text(file_content, content_type)
    chunks = chunk_text(raw_text)

    doc = Document(
        id=uuid.uuid4(),
        filename=filename,
        content_type=content_type,
        chunk_count=len(chunks),
    )
    db.add(doc)
    await db.flush()

    for i, chunk in enumerate(chunks):
        embedding = await get_embedding(chunk)
        db_chunk = DocumentChunk(
            id=uuid.uuid4(),
            document_id=doc.id,
            chunk_index=i,
            content=chunk,
            embedding=embedding,
        )
        db.add(db_chunk)

    await db.commit()
    await db.refresh(doc)
    return doc


async def retrieve_context(
    db: AsyncSession,
    query: str,
    max_results: int = None,
) -> list[dict]:
    max_results = max_results or settings.max_context_docs
    query_embedding = await get_embedding(query)

    # Cosine similarity search using pgvector
    result = await db.execute(
        text(
            """
            SELECT dc.content, dc.chunk_index, d.filename,
                   1 - (dc.embedding <=> :embedding::vector) AS score
            FROM document_chunks dc
            JOIN documents d ON dc.document_id = d.id
            ORDER BY dc.embedding <=> :embedding::vector
            LIMIT :limit
            """
        ),
        {"embedding": str(query_embedding), "limit": max_results},
    )
    rows = result.fetchall()

    return [
        {
            "content": row.content,
            "filename": row.filename,
            "chunk_index": row.chunk_index,
            "score": float(row.score),
        }
        for row in rows
    ]


def build_rag_prompt(query: str, context_chunks: list[dict]) -> str:
    context_text = "\n\n---\n\n".join(
        f"[From {c['filename']}, chunk {c['chunk_index']}]\n{c['content']}"
        for c in context_chunks
    )
    return (
        f"Use the following context to answer the question.\n\n"
        f"Context:\n{context_text}\n\n"
        f"Question: {query}"
    )
